import io
import itertools

import docuscospacy as ds
import polars as pl
import spacy
import streamlit as st
from docx import Document
from pypdf import PdfReader

st.set_page_config(page_title="DocuScope Tag Explorer", layout="wide")
st.title("DocuScope Tag Explorer")
st.write(
    "Upload a document to tag it with "
    "[docuscospacy](https://docuscospacy.readthedocs.io/en/latest/), then click a tag "
    "below to highlight every span of that category in the text."
)

# Fixed 8-hue order validated for colorblind (deuteranopia/protanopia) separation
# at full saturation, tinted toward white so black text stays legible on top.
# Text contrast held up at every tint tested down to 0.40 (worst case, violet,
# still >6:1 against black -- WCAG needs 4.5:1), so a single moderate tint (0.60)
# is used for the first 8 tags to keep as much of that separation as tinting allows.
# DocuScope documents routinely surface more than 8 categories; past 8, hues repeat
# at a second, more saturated tint band. That repetition is a real accessibility
# limit -- color can no longer carry identity alone once a hue is reused -- which is
# why every highlight also carries its tag name (hover tooltip, legend, pill label)
# as a non-color fallback. See the in-app warning shown once a document exceeds 8 tags.
BASE_HUES = [
    "#2a78d6", "#eb6834", "#1baf7a", "#eda100",
    "#e87ba4", "#008300", "#4a3aa7", "#e34948",
]
TINT_AMOUNTS = [0.60, 0.35]


def _tint(hex_color: str, amount: float) -> str:
    r, g, b = (int(hex_color[i : i + 2], 16) for i in (1, 3, 5))
    r, g, b = (round(c + (255 - c) * amount) for c in (r, g, b))
    return f"#{r:02x}{g:02x}{b:02x}"


def build_tag_colors(tags: list[str]) -> dict[str, str]:
    colors = {}
    for i, tag in enumerate(tags):
        hue = BASE_HUES[i % len(BASE_HUES)]
        tint_amount = TINT_AMOUNTS[(i // len(BASE_HUES)) % len(TINT_AMOUNTS)]
        colors[tag] = _tint(hue, tint_amount)
    return colors


@st.cache_resource(show_spinner="Loading the DocuScope language model (first run only)...")
def load_model():
    return spacy.load("en_docusco_spacy")


def extract_text(uploaded_file) -> str:
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()
    if name.endswith(".docx"):
        document = Document(io.BytesIO(data))
        return "\n".join(p.text for p in document.paragraphs)
    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)
    return data.decode("utf-8", errors="replace")


def get_tag_df(text: str, doc_id: str, nlp) -> pl.DataFrame:
    """Tag the document, memoized in this browser session only.

    Deliberately not st.cache_data: that cache is process-wide and outlives any
    one session, which would keep uploaded text in server memory after the tab
    that uploaded it is gone. Session state is discarded when the session ends
    (tab closed/refreshed), and nothing here ever touches disk.
    """
    cache_key = (doc_id, len(text), hash(text))
    if st.session_state.get("_cache_key") != cache_key:
        with st.spinner("Tagging the document with DocuScope..."):
            corpus = pl.DataFrame({"doc_id": [doc_id], "text": [text]})
            tokens = ds.docuscope_parse(corpus, nlp)
            st.session_state["_tag_df"] = ds.tag_ruler(tokens, doc_id=doc_id, count_by="ds")
        st.session_state["_cache_key"] = cache_key
    return st.session_state["_tag_df"]


def build_spans(tag_df: pl.DataFrame) -> list[dict]:
    spans = []
    for _tag_id, rows in itertools.groupby(tag_df.iter_rows(named=True), key=lambda r: r["tag_id"]):
        rows = list(rows)
        spans.append({"text": "".join(r["token"] for r in rows), "tag": rows[0]["tag"]})
    return spans


def render_legend(selected_tags: list[str], colors: dict[str, str]) -> str:
    badges = [
        f'<span style="background-color:{colors[tag]};border-radius:3px;'
        f'padding:2px 8px;margin:0 6px 6px 0;display:inline-block;font-size:0.85rem;">'
        f"{tag}</span>"
        for tag in selected_tags
    ]
    return "".join(badges)


def render_html(spans: list[dict], colors: dict[str, str], selected: set[str]) -> str:
    escape = lambda t: t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    parts = []
    for span in spans:
        text = escape(span["text"])
        if span["tag"] in selected:
            color = colors[span["tag"]]
            parts.append(
                f'<mark title="{span["tag"]}" '
                f'style="background-color:{color};border-radius:3px;padding:0 1px;">'
                f"{text}</mark>"
            )
        else:
            parts.append(text)
    return "".join(parts)


uploaded_file = st.file_uploader("Upload a document", type=["txt", "docx", "pdf"])
st.caption(
    "Your file is processed in memory for this browser session only — it is never "
    "written to disk, and everything is discarded when you remove the file, refresh, "
    "or close this tab."
)

if uploaded_file is None:
    # No active upload: drop anything left over from a previous file in this session.
    st.session_state.pop("_cache_key", None)
    st.session_state.pop("_tag_df", None)
else:
    text = extract_text(uploaded_file)
    if not text.strip():
        st.warning("No text could be extracted from this file.")
    else:
        nlp = load_model()
        tag_df = get_tag_df(text, uploaded_file.name, nlp)
        spans = build_spans(tag_df)

        all_tags = sorted(
            tag_df.filter(pl.col("tag") != "Untagged")["tag"].unique().to_list()
        )
        colors = build_tag_colors(all_tags)

        if not all_tags:
            st.info("DocuScope didn't find any tagged spans in this document.")
            selected: set[str] = set()
        else:
            st.write("**DocuScope tags** (click to highlight):")
            selected = set(
                st.pills(
                    "DocuScope tags",
                    all_tags,
                    selection_mode="multi",
                    label_visibility="collapsed",
                    key="tag_pills",
                )
                or []
            )
            st.caption(
                "Colors use a fixed, colorblind-safe hue order (checked for deuteranopia "
                "and protanopia). Don't rely on color alone: each highlight shows its tag "
                "name on hover, and it appears in the legend above the text whenever 2+ "
                "tags are selected."
            )
            if len(all_tags) > len(BASE_HUES):
                st.warning(
                    f"This document has {len(all_tags)} tagged categories, more than the "
                    f"{len(BASE_HUES)} fully distinct colors available. Categories past "
                    f"the first {len(BASE_HUES)} reuse a hue at a different shade, so "
                    "color alone won't reliably distinguish them — use the tag name in "
                    "the tooltip or legend instead."
                )

        # A single, height-reserved legend row (populated only past one selection)
        # so toggling pills never shifts anything above or below it.
        legend_html = render_legend(sorted(selected), colors) if len(selected) > 1 else ""
        st.markdown(
            f'<div style="min-height:1.9rem;margin:0.4rem 0;">{legend_html}</div>',
            unsafe_allow_html=True,
        )

        st.divider()
        html = render_html(spans, colors, selected)
        st.markdown(
            f'<div style="white-space:pre-wrap;line-height:1.6;">{html}</div>',
            unsafe_allow_html=True,
        )
