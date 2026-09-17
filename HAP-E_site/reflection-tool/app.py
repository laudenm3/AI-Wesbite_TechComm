# Reflection worksheet — a few free-response text boxes that get concatenated
# into a single .docx the user can download. Meant to be deployed (e.g. to
# Streamlit Community Cloud) and embedded on the main site via <iframe>.

import io

import streamlit as st
from docx import Document

st.set_page_config(page_title="Reflection Worksheet", layout="centered")

# Hide Streamlit's own chrome (menu/footer/header) so the app blends into the
# page it's embedded in via iframe.
st.markdown(
    """
    <style>
    #MainMenu, footer, header {visibility: hidden;}
    .block-container {padding-top: 1.5rem; padding-bottom: 2rem;}
    </style>
    """,
    unsafe_allow_html=True,
)

st.title("Reflection Worksheet")
st.caption("Answer the prompts below, then click the button to download your responses as a Word document.")

# --- QUESTIONS — edit these prompts to match your assignment ---
QUESTIONS = [
    "Question 1: ...",
    "Question 2: ...",
    "Question 3: ...",
    "Question 4: ...",
]

answers = [
    st.text_area(question, key=f"answer_{i}", height=120)
    for i, question in enumerate(QUESTIONS)
]

doc = Document()
doc.add_heading("Reflection Worksheet Responses", level=1)
for question, answer in zip(QUESTIONS, answers):
    doc.add_heading(question, level=2)
    doc.add_paragraph(answer.strip() or "(No response provided.)")

buffer = io.BytesIO()
doc.save(buffer)

st.download_button(
    label="⬇️ Save answers as Word document",
    data=buffer.getvalue(),
    file_name="reflection_responses.docx",
    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    use_container_width=True,
)
